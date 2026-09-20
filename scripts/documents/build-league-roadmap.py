"""Generate the roadmap and parallel-development agreement as one Word document."""

import argparse
from datetime import datetime, timezone
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DIRECTORY = ROOT / "docs/development/league"
SOURCES = ["roadmap-20260919.md", "parallel-development-20260919.md"]
OUTPUT = DIRECTORY / "数字攻防联赛开发路线与并行规范V0.2.docx"


def font(style, face, size, color="000000"):
    style.font.name = face
    style.font.size = Pt(size)
    style.font.bold = False
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    for attribute in list(rpr.rFonts.attrib):
        if "theme" in attribute.lower():
            del rpr.rFonts.attrib[attribute]
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        rpr.rFonts.set(qn("w:" + name), face)
    for tag, values in (("bCs", {"val": "0"}), ("szCs", {"val": str(int(size * 2))}),
                        ("lang", {"val": "zh-CN", "eastAsia": "zh-CN"})):
        node = rpr.find(qn("w:" + tag))
        if node is None:
            node = OxmlElement("w:" + tag)
            rpr.append(node)
        for name, value in values.items():
            node.set(qn("w:" + name), value)


def rich(paragraph, text, source):
    for part in re.split(r"(`[^`]+`|\[[^\]]+\]\([^)]+\))", text):
        match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if match:
            label, target = match.groups()
            if not target.startswith(("https://", "http://")):
                relative = (source.parent / target).resolve().relative_to(ROOT).as_posix()
                target = "https://github.com/Y3X1L2/newGZCTF/blob/codex/league-roadmap-20260919/" + relative
            relationship = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), relationship)
            run = OxmlElement("w:r")
            node = OxmlElement("w:t")
            node.text = label
            run.append(node)
            link.append(run)
            paragraph._p.append(link)
        else:
            run = paragraph.add_run(part.strip("`") if part.startswith("`") else part)
            if part.startswith("`"):
                run.font.name = "Consolas"
                run.font.size = Pt(10)


def add_table(doc, lines, source):
    rows = [[v.strip() for v in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", v) for v in row)]
    if len(rows[0]) == 2:
        widths = [4.2, 12.2]
    else:
        widths = {
            "证据": [2.8, 7.0, 6.6], "失败测试": [6.5, 4.5, 5.4],
            "原规划安排": [4.0, 5.5, 6.9], "人员": [1.5, 8.0, 6.9],
            "任务": [4.2, 4.8, 7.4], "契约": [3.2, 7.0, 6.2],
            "阶段与工作": [6.0, 2.3, 8.1], "事项": [5.0, 4.0, 7.4],
            "目录范围": [8.2, 5.0, 3.2], "共享位置": [7.5, 2.6, 6.3],
            "批次": [1.3, 8.0, 7.1], "环境": [3.8, 6.0, 6.6],
            "情况": [5.5, 4.2, 6.7], "PR 示例": [5.1, 5.3, 6.0],
            "负责人": [1.8, 8.0, 6.6], "时间": [2.6, 7.0, 6.8],
        }.get(rows[0][0], [4.0, 5.0, 7.4])
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + name)
        for key, value in {"val": "single", "sz": "4", "color": "D1D5DB"}.items():
            node.set(qn("w:" + key), value)
        borders.append(node)
    table._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for name, value in (("top", "80"), ("bottom", "80"), ("left", "100"), ("right", "100")):
        node = OxmlElement("w:" + name)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    table._tbl.tblPr.append(margins)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if not index:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.style = "Roadmap Table" if index else "Roadmap Table Header"
            rich(paragraph, value, source)
            if not index:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EFF5")
                cell._tc.get_or_add_tcPr().append(shade)
    gap = doc.add_paragraph()
    gap.paragraph_format.line_spacing = Pt(1)
    gap.paragraph_format.space_after = Pt(4)
    gap.add_run().font.size = Pt(1)


def build(sources=SOURCES, output=OUTPUT,
          title="数字攻防联赛开发路线与并行规范",
          subject="三阶段评估、开发任务、责任边界与分支合并安排",
          date="2026-09-19"):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.9)
    section.left_margin = section.right_margin = Cm(2.3)
    section.footer_distance = Cm(0.8)
    normal = doc.styles["Normal"]
    font(normal, "SimSun", 12)
    normal.paragraph_format.line_spacing = Pt(18)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[name]
        font(style, "SimHei", size)
        link = style.element.find(qn("w:link"))
        if link is not None:
            style.element.remove(link)
        style.paragraph_format.space_before = Pt(8 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    for style in doc.styles:
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    for name, face, size in (("Roadmap Table", "SimSun", 10.5), ("Roadmap Table Header", "SimHei", 10.5),
                              ("Roadmap Code", "Consolas", 9)):
        style = doc.styles.add_style(name, 1)
        style.base_style = normal
        font(style, face, size)
        if name == "Roadmap Code":
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.paragraph_format.line_spacing = Pt(14 if name == "Roadmap Code" else 16)
        style.paragraph_format.space_after = Pt(2)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("第 ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页").font.size = Pt(9)
    for source_index, filename in enumerate(sources):
        source = DIRECTORY / filename
        lines = source.read_text(encoding="utf-8").splitlines()
        i = 0
        pending_break = bool(source_index)
        code = False
        while i < len(lines):
            line = lines[i].strip()
            i += 1
            if line.startswith("```"):
                code = not code
                continue
            if not line:
                continue
            if line == "<!-- pagebreak -->":
                pending_break = True
                continue
            if code:
                doc.add_paragraph(line, style="Roadmap Code")
            elif line.startswith("|"):
                group = [line]
                while i < len(lines) and lines[i].strip().startswith("|"):
                    group.append(lines[i].strip())
                    i += 1
                add_table(doc, group, source)
            else:
                match = re.match(r"^(#{1,3}) (.*)$", line)
                if match:
                    style = {1: "Title", 2: "Heading 1", 3: "Heading 2"}[len(match[1])]
                    paragraph = doc.add_paragraph(match[2], style=style)
                else:
                    paragraph = doc.add_paragraph()
                    rich(paragraph, line, source)
                    if re.match(r"\d+\. ", line):
                        paragraph.paragraph_format.left_indent = Cm(0.45)
                        paragraph.paragraph_format.first_line_indent = Cm(-0.45)
                if pending_break:
                    paragraph.paragraph_format.page_break_before = True
                    pending_break = False
    props = doc.core_properties
    props.title = title
    props.author = "YINYU"
    props.subject = subject
    props.created = props.modified = datetime.fromisoformat(date).replace(tzinfo=timezone.utc)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", action="append", help="Markdown filename under docs/development/league")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--title", default="数字攻防联赛开发路线与并行规范")
    parser.add_argument("--subject", default="三阶段评估、开发任务、责任边界与分支合并安排")
    parser.add_argument("--date", default="2026-09-19")
    args = parser.parse_args()
    build(args.source or SOURCES, args.output, args.title, args.subject, args.date)
