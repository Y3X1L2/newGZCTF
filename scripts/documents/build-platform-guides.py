"""Build the platform guides from their reviewed Markdown sources.

Run with the document runtime's Python. Screenshot instructions stay in the
Word files until the editor replaces them with the requested screenshots.
"""

from datetime import datetime, timezone
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
GUIDES = ROOT / "docs" / "user-guides"
OUTPUTS = {
    "platform-product-overview.md": "隐喻平台产品说明V0.1.docx",
    "platform-user-manual.md": "隐喻平台使用手册V0.1.docx",
}


def font(style, name, size, bold=False, color="000000"):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def rich(paragraph, text):
    for part in re.split(r"(`[^`]+`)", text):
        run = paragraph.add_run(part.strip("`") if part.startswith("`") else part)
        if part.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)


def table(document, lines):
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    t = document.add_table(rows=0, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [3.0, 7.4, 6.0] if len(rows[0]) == 3 else [4.2, 12.2]
    if len(rows[0]) == 3 and rows[0][1] == "截图编号":
        widths = [3.6, 3.0, 9.8]
    for column, width in zip(t.columns, widths):
        column.width = Cm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        for key, value in {"val": "single", "sz": "4", "color": "D1D5DB"}.items():
            edge.set(qn("w:" + key), value)
        borders.append(edge)
    pr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", "100"), ("bottom", "100"), ("left", "120"), ("right", "120")):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), value)
        e.set(qn("w:type"), "dxa")
        margins.append(e)
    pr.append(margins)
    for index, values in enumerate(rows):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for col, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = document.styles["Guide Table"]
            rich(p, value)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EFF5")
                cell._tc.get_or_add_tcPr().append(shade)
                for run in p.runs:
                    run.bold = True
    gap = document.add_paragraph()
    gap.paragraph_format.space_after = Pt(4)
    gap.paragraph_format.line_spacing = Pt(1)
    gap.add_run().font.size = Pt(1)


def build(source, output):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2.3), Cm(2.3)
    section.footer_distance = Cm(0.8)
    normal = document.styles["Normal"]
    font(normal, "微软雅黑", 12)
    normal.paragraph_format.line_spacing = Pt(18)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 24), ("Heading 1", 17), ("Heading 2", 13)):
        style = document.styles[name]
        font(style, "微软雅黑", size, True)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    for style in document.styles:
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    for name, size, color in (("Guide Shot", 10.5, "4B5563"), ("Guide Table", 11, "000000")):
        style = document.styles.add_style(name, 1)
        style.base_style = normal
        font(style, "微软雅黑", size, color=color)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = Pt(15.5 if name == "Guide Shot" else 16.5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("第 ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页").font.size = Pt(9)
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    in_shot = False
    pending_break = False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line == "<!-- pagebreak -->":
            pending_break = True
            continue
        elif line.startswith("::: screenshot "):
            in_shot = True
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
            p.add_run("待补截图 " + line[len("::: screenshot "):]).bold = True
        elif line == ":::":
            in_shot = False
        elif in_shot:
            p = document.add_paragraph(style="Guide Shot")
            rich(p, line)
            p.paragraph_format.keep_with_next = not line.startswith("图注：")
        elif line.startswith("| "):
            group = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                group.append(lines[i].strip())
                i += 1
            table(document, group)
        elif line.startswith("# "):
            document.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            p = document.add_paragraph(line[3:], style="Heading 1")
            if pending_break:
                p.paragraph_format.page_break_before = True
                pending_break = False
        elif line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
        else:
            p = document.add_paragraph()
            rich(p, line)
            if re.match(r"\d+\. ", line):
                p.paragraph_format.left_indent = Cm(0.45)
                p.paragraph_format.first_line_indent = Cm(-0.45)
    props = document.core_properties
    props.title = next(line[2:] for line in lines if line.startswith("# "))
    props.subject = "平台功能和操作说明，附截图补充指引"
    props.author = "YINYU"
    props.keywords = "隐喻平台,培训,练习,赛事,使用手册"
    props.created = props.modified = datetime(2026, 9, 14, tzinfo=timezone.utc)
    document.save(output)
    print(output)


if __name__ == "__main__":
    for name, output_name in OUTPUTS.items():
        build(GUIDES / name, GUIDES / output_name)
